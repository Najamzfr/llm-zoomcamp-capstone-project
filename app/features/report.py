"""Literature-review report (F12): one structured llm_json call over an answer's sources.

Rendered in a panel and downloadable as Markdown and .docx (python-docx). References list only the
papers the review cites; nothing about how the sources were found.
"""

import io
import re

import streamlit as st

from research_assistant.presentation import link_citations

from .export import reference_line, slug
from .llm_extras import LitReview, LLMExtrasError, llm_json
from .records import sources_context

PROMPT = (
    "Write a short literature review of these sources about the question. Return: a title; an overview "
    "(3-4 sentences); 2 to 4 themes, each with a title, a 2-3 sentence summary and `src`; a methods "
    "comparison (method, core idea, strengths, limitations, `src`) for the approaches the sources describe; "
    "a timeline of key ideas ordered by year as given in the sources (`when` is a year or 'Earlier work'); "
    "2 to 4 open problems; and `cited`: every source number you used. Cite inline with [n] in the overview "
    "and summaries, using the numbers from the context."
)


def generate(record):
    data = llm_json(PROMPT, LitReview, sources_context(record), feature="report", answer_id=record["id"])
    valid = {source["n"] for source in record.get("sources", [])}
    used = set(data.get("cited") or [])
    for section in ("themes", "methods", "timeline"):
        for item in data.get(section, []):
            item["src"] = [n for n in item.get("src") or [] if n in valid]
            used.update(item["src"])
    for text in [data["overview"]] + [theme["summary"] for theme in data["themes"]]:
        used.update(int(n) for n in re.findall(r"\[(\d{1,2})\]", text))
    data["cited"] = sorted(n for n in used if n in valid)
    return data


def _cites(src, text=""):
    """Bracketed source numbers, skipping any the text already cites inline."""
    present = {int(n) for n in re.findall(r"\[(\d{1,2})\]", text)}
    return "".join(f"[{n}]" for n in src if n not in present)


def references(record, data):
    by_n = {source["n"]: source for source in record.get("sources", [])}
    return [f"[{n}] {reference_line(by_n[n])}" for n in data["cited"] if n in by_n]


def to_markdown(record, data):
    lines = [f"# {data['title']}", "", f"*Question: {record['question']}*", "", "## Overview", "", data["overview"], "",
             "## Themes", ""]
    for theme in data["themes"]:
        lines += [f"### {theme['title']}", "", f"{theme['summary']} {_cites(theme['src'], theme['summary'])}".strip(), ""]
    if data["methods"]:
        lines += ["## Methods compared", "", "| Method | Core idea | Strengths | Limitations | Sources |", "|---|---|---|---|---|"]
        for row in data["methods"]:
            cells = [row["method"], row["idea"], row["strengths"], row["limitations"], _cites(row["src"])]
            lines.append("| " + " | ".join(cell.replace("|", "/") for cell in cells) + " |")
        lines.append("")
    if data["timeline"]:
        lines += ["## Timeline of key ideas", ""]
        lines += [f"- **{item['when']}** — {item['idea']} {_cites(item['src'], item['idea'])}".rstrip() for item in data["timeline"]]
        lines.append("")
    if data["open_problems"]:
        lines += ["## Open problems", ""] + [f"- {problem}" for problem in data["open_problems"]] + [""]
    lines += ["## References", ""] + references(record, data)
    return "\n".join(lines).strip() + "\n"


def to_docx(record, data):
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(data["title"], level=0)
    document.add_paragraph(f"Question: {record['question']}").italic = True
    document.add_heading("Overview", level=1)
    document.add_paragraph(data["overview"])
    document.add_heading("Themes", level=1)
    for theme in data["themes"]:
        document.add_heading(theme["title"], level=2)
        document.add_paragraph(f"{theme['summary']} {_cites(theme['src'], theme['summary'])}".strip())
    if data["methods"]:
        document.add_heading("Methods compared", level=1)
        table = document.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for cell, title in zip(table.rows[0].cells, ("Method", "Core idea", "Strengths", "Limitations", "Sources")):
            cell.text = title
        for row in data["methods"]:
            cells = table.add_row().cells
            for cell, value in zip(cells, (row["method"], row["idea"], row["strengths"], row["limitations"], _cites(row["src"]))):
                cell.text = value
    if data["timeline"]:
        document.add_heading("Timeline of key ideas", level=1)
        for item in data["timeline"]:
            document.add_paragraph(f"{item['when']} — {item['idea']} {_cites(item['src'], item['idea'])}".strip(), style="List Bullet")
    if data["open_problems"]:
        document.add_heading("Open problems", level=1)
        for problem in data["open_problems"]:
            document.add_paragraph(problem, style="List Bullet")
    document.add_heading("References", level=1)
    for line in references(record, data):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_panel(record, mid, prefix):
    try:
        with st.spinner("Writing a literature review…"):
            data = generate(record)
    except LLMExtrasError as error:
        st.error(str(error))
        return
    count = len(record.get("sources", []))
    st.html("<div class='panel-h'><b>Literature review</b></div>")
    body = to_markdown(record, data).split("## References", 1)[0]
    body = body.replace(f"# {data['title']}", f"#### {data['title']}", 1).replace("\n## ", "\n##### ").replace("\n### ", "\n###### ")
    with st.container(key=f"prose_report_{mid}"):
        st.markdown(link_citations(body, prefix, count), unsafe_allow_html=True)
    st.caption("References: " + "; ".join(f"[{n}]" for n in data["cited"]) + " — open them with Sources.")
    base = slug(record["question"], 50) + "-review"
    with st.container(horizontal=True, key=f"report_{mid}_tools"):
        st.download_button("Download review (.md)", to_markdown(record, data), file_name=f"{base}.md",
                           mime="text/markdown", key=f"report_{mid}_md")
        st.download_button("Download review (.docx)", to_docx(record, data), file_name=f"{base}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           key=f"report_{mid}_docx")
