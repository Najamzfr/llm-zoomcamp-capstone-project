"""Phase 3 unit tests: PDF uploads + fusion, literature review export, read-aloud text, collections, feedback."""

import os

os.environ["LANGSMITH_TRACING"] = "false"

import io
import json
import unittest
from unittest import mock

import numpy as np
import streamlit as st

from app.features import collections, feedback, llm_extras, report, share, tts, uploads
from app.storage import local
from research_assistant.agent import AdaptiveResearchAgent, fuse_rankings
from tests.test_extras import record
from tests.test_extras_phase2 import reset_storage

REPORT_REPLY = json.dumps({
    "title": "Efficient fine-tuning",
    "overview": "Low-rank updates [1] and adapters [2] cut training cost.",
    "themes": [{"title": "Parameter efficiency", "summary": "Both train few weights [1].", "src": [1, 2]}],
    "methods": [{"method": "LoRA", "idea": "Low-rank deltas", "strengths": "No latency", "limitations": "Rank choice", "src": [1]}],
    "timeline": [{"when": "2019", "idea": "Adapters", "src": [2]}, {"when": "2021", "idea": "LoRA", "src": [1, 7]}],
    "open_problems": ["Choosing ranks automatically"],
    "cited": [1, 2, 5],
})


def make_pdf(pages, text="Retrieval augmented generation combines search with language models. "):
    import fitz

    document = fitz.open()
    for number in range(pages):
        page = document.new_page()
        if text:
            page.insert_textbox(fitz.Rect(40, 40, 560, 800), f"Page {number + 1}. " + text * 12, fontsize=9)
    data = document.tobytes()
    document.close()
    return data


class FakeEmbedder:
    """Deterministic bag-of-words vectors so tests don't load the ONNX model."""

    words = ["retrieval", "generation", "search", "language", "diffusion", "robot", "page", "models"]

    def _vector(self, text):
        lowered = text.lower()
        vector = np.array([lowered.count(word) for word in self.words], dtype="float32") + 0.01
        return vector / np.linalg.norm(vector)

    def embed_documents(self, texts, batch_size=4):
        return np.stack([self._vector(text) for text in texts])

    def embed_query(self, text):
        return self._vector(text)


class UploadTests(unittest.TestCase):
    def setUp(self):
        reset_storage()

    def test_extract_and_chunk_with_overlap(self):
        pages = uploads.extract_pages(make_pdf(3))
        self.assertEqual([number for number, _ in pages], [1, 2, 3])
        chunks = uploads.chunk_pages(pages, "paper.pdf", tokens=120, overlap=20)
        self.assertGreater(len(chunks), 3)
        import tiktoken

        encoding = tiktoken.get_encoding("o200k_base")
        sizes = [len(encoding.encode(chunk["text"])) for chunk in chunks]
        self.assertLessEqual(max(sizes), 122)
        self.assertEqual(chunks[0]["page"], 1)
        self.assertEqual(chunks[-1]["page"], 3)
        self.assertIn(chunks[0]["text"][-30:], chunks[1]["text"][:200])  # 20-token overlap

    def test_friendly_errors(self):
        with self.assertRaisesRegex(uploads.UploadError, "no selectable text"):
            uploads.extract_pages(make_pdf(2, text=""))
        with self.assertRaisesRegex(uploads.UploadError, "limit is about 60"):
            uploads.extract_pages(make_pdf(61, text="x "))
        with self.assertRaisesRegex(uploads.UploadError, "isn't a PDF"):
            uploads.extract_pages(b"hello")

    def test_link_fetch_blocks_private_addresses(self):
        for url in ("http://localhost:8000/x.pdf", "http://127.0.0.1/x.pdf", "file:///etc/passwd", "ftp://x/y.pdf"):
            with self.assertRaises(uploads.UploadError):
                uploads.fetch_pdf(url)

    def test_session_index_search_and_clear(self):
        with mock.patch.object(uploads, "_embedder", return_value=FakeEmbedder()):
            self.assertEqual(uploads.add("my-draft.pdf", make_pdf(10)), 10)
            hits = uploads.searcher()("retrieval generation", 2)
            self.assertEqual(len(hits), 2)
            self.assertTrue(all(hit["upload"] and hit["title"] == "my-draft" for hit in hits))
            self.assertTrue(hits[0]["paper_id"].startswith("upload:my-draft.pdf:c"))
            with self.assertRaisesRegex(uploads.UploadError, "already added"):
                uploads.add("my-draft.pdf", make_pdf(1))
            uploads.remove("my-draft.pdf")
        self.assertEqual(uploads.files(), [])
        self.assertIsNone(uploads.searcher())


class FusionTests(unittest.TestCase):
    def test_rrf_interleaves_and_renumbers(self):
        papers = [{"paper_id": f"p{i}", "rank": i} for i in range(1, 6)]
        extra = [{"paper_id": "upload:a.pdf:c1", "upload": True}, {"paper_id": "upload:a.pdf:c2", "upload": True}]
        fused = fuse_rankings(papers, extra, 7)
        self.assertEqual([paper["rank"] for paper in fused], list(range(1, 8)))
        self.assertEqual({paper["paper_id"] for paper in fused[:4]}, {"p1", "upload:a.pdf:c1", "p2", "upload:a.pdf:c2"})

    def test_agent_never_downloads_uploads(self):
        agent = AdaptiveResearchAgent.__new__(AdaptiveResearchAgent)
        papers = [{"paper_id": "2106.09685"}, {"paper_id": "upload:a.pdf:c1", "upload": True}]
        decision = {"selected_paper_ids": ["upload:a.pdf:c1", "2106.09685"]}
        self.assertEqual(agent._selected_ids(decision, papers), ["2106.09685"])
        import tiktoken

        agent.encoding = tiktoken.get_encoding("o200k_base")
        context = agent._with_uploads("[1] Paper\ntext", [{"rank": 2, "title": "a", "page": 3, "text": "My result."}])
        self.assertIn("[2] a (your upload; page: 3)\nMy result.", context)


class ReportTests(unittest.TestCase):
    def setUp(self):
        llm_extras._cached.clear()

    def test_markdown_and_docx(self):
        with mock.patch.object(llm_extras, "_complete", return_value=REPORT_REPLY):
            data = report.generate(record())
        self.assertEqual(data["cited"], [1, 2])  # 5 and 7 don't exist
        markdown = report.to_markdown(record(), data)
        for heading in ("## Overview", "## Themes", "## Methods compared", "## Timeline of key ideas", "## Open problems",
                        "## References"):
            self.assertIn(heading, markdown)
        self.assertIn("| LoRA | Low-rank deltas | No latency | Rank choice | [1] |", markdown)
        self.assertIn("[2] Parameter-Efficient Transfer Learning for NLP", markdown)
        self.assertIn("Both train few weights [1]. [2]", markdown)  # [1] already inline, not repeated
        self.assertNotIn("cs.", markdown)
        from docx import Document

        document = Document(io.BytesIO(report.to_docx(record(), data)))
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("Efficient fine-tuning", texts)
        self.assertIn("Methods compared", texts)
        self.assertEqual(len(document.tables[0].rows), 2)


class SpeechTests(unittest.TestCase):
    def test_short_answer_first_without_marks(self):
        text = tts.speech_text("Short answer: LoRA is cheap [1].\n\n#### Why\n| a | b |\n|---|---|\n- It trains **few** weights [2][3].")
        self.assertTrue(text.startswith("LoRA is cheap."))
        for mark in ("[", "|", "#", "*"):
            self.assertNotIn(mark, text)
        self.assertIn("It trains few weights.", text)


class CollectionTests(unittest.TestCase):
    def setUp(self):
        reset_storage()

    def test_save_rename_export_and_delete(self):
        rec = record()
        rec["sources"].append({"n": 3, "title": "my-draft", "upload": True, "excerpt": "secret", "evidence": "secret"})
        self.assertTrue(collections.save("  Thesis   reading ", rec))
        self.assertFalse(collections.save("Thesis reading", rec))
        stored = collections.all_collections()["Thesis reading"][0]
        self.assertNotIn("secret", json.dumps(stored))
        with self.assertRaises(ValueError):
            collections.save("   ", rec)
        self.assertEqual(collections.rename("Thesis reading", "x" * 60), "x" * 40)
        name = "x" * 40
        items = collections.all_collections()[name]
        markdown = collections.collection_markdown(name, items)
        self.assertTrue(markdown.startswith(f"# {name}"))
        self.assertIn("## Compare LoRA and adapters?", markdown)
        self.assertEqual(collections.collection_bibtex(items + items).count("@misc{"), 2)
        collections.remove(name, rec["id"])
        self.assertEqual(collections.all_collections()[name], [])
        collections.delete(name)
        self.assertEqual(collections.all_collections(), {})

    def test_history_and_share_drop_uploaded_text(self):
        from app.features import history

        rec = record()
        rec["sources"].append({"n": 3, "title": "my-draft", "upload": True, "excerpt": "secret", "evidence": "secret"})
        history.add(rec)
        self.assertNotIn("secret", json.dumps(local.get(local.HISTORY)))
        self.assertNotIn("secret", json.dumps(share.snapshot(rec)))


class FeedbackTests(unittest.TestCase):
    def test_rows_have_no_identifiers(self):
        reset_storage()
        share._memory_backend.clear()
        with mock.patch.dict(os.environ, {"SHARE_BACKEND": "memory"}):
            self.assertTrue(feedback.submit(record(), "down", "Missing source", " needs more " * 100))
            saved = share.backend().feedback[-1]
        self.assertEqual(set(saved), {"answer_id", "question", "rating", "reason", "text", "created_at"})
        self.assertEqual(saved["reason"], "Missing source")
        self.assertLessEqual(len(saved["text"]), feedback.MAX_TEXT)
        with mock.patch.dict(os.environ, {"SHARE_BACKEND": "supabase", "SUPABASE_URL": ""}):
            self.assertFalse(feedback.submit(record(), "up"))


if __name__ == "__main__":
    unittest.main()
